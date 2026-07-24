"""Generate the final three-slide Sentinel presentation from project facts."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Sentinel_Final_Deck.pptx"
GREEN = RGBColor(13, 48, 40)
EMERALD = RGBColor(4, 120, 87)
LIME = RGBColor(190, 242, 100)
INK = RGBColor(28, 35, 32)
MUTED = RGBColor(91, 103, 97)
PAPER = RGBColor(248, 248, 244)
WHITE = RGBColor(255, 255, 255)


def textbox(slide, x, y, w, h, text, size=20, color=INK, bold=False, font="Aptos", align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    run = paragraph.runs[0]
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def box(slide, x, y, w, h, fill=WHITE, radius=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = RGBColor(220, 224, 219)
    return shape


def base_slide(prs, number, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid(); background.fore_color.rgb = PAPER
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.16), prs.slide_height)
    accent.fill.solid(); accent.fill.fore_color.rgb = EMERALD; accent.line.fill.background()
    textbox(slide, 0.55, 0.28, 1.5, 0.25, f"SENTINEL  /  0{number}", 10, EMERALD, True)
    textbox(slide, 0.55, 0.62, 12.1, 0.55, title, 28, INK, True)
    textbox(slide, 0.57, 1.16, 11.9, 0.35, subtitle, 13, MUTED)
    return slide


def add_bullets(shape, items, size=15, color=INK):
    frame = shape.text_frame
    frame.clear(); frame.word_wrap = True
    for index, item in enumerate(items):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = item; p.level = 0; p.space_after = Pt(8)
        p.font.name = "Aptos"; p.font.size = Pt(size); p.font.color.rgb = color
        p.text = f"•  {item}"


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = base_slide(prs, 1, "From assignment to shipped product", "A local RAG experiment became a complete, auditable cybersecurity workspace.")
    box(slide, 0.55, 1.75, 4.0, 3.75)
    textbox(slide, 0.82, 2.0, 3.4, 0.35, "WHAT WAS ASKED", 12, EMERALD, True)
    asked = slide.shapes.add_textbox(Inches(0.82), Inches(2.48), Inches(3.35), Inches(2.65))
    add_bullets(asked, ["Local CIS Controls RAG", "Connected chat frontend", "Streaming + citations", "Persistence + feedback"], 16)
    box(slide, 4.72, 1.75, 4.05, 3.75, RGBColor(236, 253, 245))
    textbox(slide, 5.0, 2.0, 3.5, 0.35, "WHAT WAS SHIPPED", 12, EMERALD, True)
    shipped = slide.shapes.add_textbox(Inches(5.0), Inches(2.48), Inches(3.4), Inches(2.65))
    add_bullets(shipped, ["Grounded SSE answers", "MongoDB threads + versions", "Interactive evidence + abstention", "Authenticated feedback loop"], 16)
    box(slide, 8.95, 1.75, 3.82, 3.75, GREEN)
    textbox(slide, 9.25, 2.0, 3.2, 0.35, "MEASURED PROOF", 12, LIME, True)
    textbox(slide, 9.25, 2.55, 3.0, 0.55, "332", 34, WHITE, True)
    textbox(slide, 10.25, 2.69, 2.0, 0.35, "vectors", 14, RGBColor(190, 210, 202))
    textbox(slide, 9.25, 3.35, 3.0, 0.55, "100%", 34, WHITE, True)
    textbox(slide, 10.55, 3.49, 1.5, 0.35, "Hit@5", 14, RGBColor(190, 210, 202))
    textbox(slide, 9.25, 4.15, 3.0, 0.55, "97.83%", 31, WHITE, True)
    textbox(slide, 10.85, 4.28, 1.4, 0.4, "pass", 14, RGBColor(190, 210, 202))
    textbox(slide, 0.58, 5.85, 12.0, 0.5, "React · .NET · FastAPI · Unstructured · BGE · Weaviate · Ollama/Qwen · MongoDB · DeepEval", 15, GREEN, True, align=PP_ALIGN.CENTER)

    slide = base_slide(prs, 2, "Architecture and grounded AI flow", "One boundary separates user experience, application security, and ML experimentation.")
    layers = [(0.7, "REACT", "Chat · Markdown · SSE", RGBColor(236, 253, 245)), (3.55, ".NET", "Auth · ownership · persistence", RGBColor(220, 252, 231)), (6.4, "FASTAPI RAG", "Retrieve · rerank · generate", RGBColor(209, 250, 229)), (9.25, "LOCAL AI", "Weaviate · BGE · Qwen", GREEN)]
    for x, title, sub, fill in layers:
        box(slide, x, 2.0, 2.35, 1.25, fill)
        color = WHITE if fill == GREEN else GREEN
        textbox(slide, x + 0.15, 2.23, 2.05, 0.35, title, 17, color, True, align=PP_ALIGN.CENTER)
        textbox(slide, x + 0.15, 2.68, 2.05, 0.3, sub, 10, color, align=PP_ALIGN.CENTER)
    for x in (3.14, 5.99, 8.84):
        textbox(slide, x, 2.38, 0.4, 0.4, "→", 24, EMERALD, True, align=PP_ALIGN.CENTER)
    box(slide, 3.55, 3.65, 5.2, 2.2, WHITE)
    textbox(slide, 3.88, 3.94, 4.55, 0.35, "JUSTIFIED DECISION", 12, EMERALD, True)
    textbox(slide, 3.88, 4.4, 4.55, 1.05, ".NET is the security and persistence boundary. The browser cannot reach the model, vector store, or database directly; Python remains focused on the RAG workflow.", 16, INK)
    textbox(slide, 0.9, 6.3, 11.6, 0.38, "Insufficient evidence → abstain. Strong evidence → Qwen answer + clickable page citations.", 16, GREEN, True, align=PP_ALIGN.CENTER)

    slide = base_slide(prs, 3, "Lessons learned and business value", "The feedback loop turns individual failures into measurable system improvements.")
    box(slide, 0.65, 1.75, 5.85, 4.45)
    textbox(slide, 0.95, 2.02, 5.1, 0.35, "LESSONS LEARNED", 12, EMERALD, True)
    lessons = slide.shapes.add_textbox(Inches(0.95), Inches(2.52), Inches(5.1), Inches(3.2))
    add_bullets(lessons, ["Retrieval quality matters more than prompt size.", "Streaming improves perception; persistence makes a product.", "Golden data, LLM judging, and human feedback serve different roles.", "AI coding agents need repository-specific constraints and verification."], 16)
    box(slide, 6.8, 1.75, 5.85, 4.45, RGBColor(236, 253, 245))
    textbox(slide, 7.1, 2.02, 5.1, 0.35, "BUSINESS VALUE", 12, EMERALD, True)
    value = slide.shapes.add_textbox(Inches(7.1), Inches(2.52), Inches(5.1), Inches(3.2))
    add_bullets(value, ["Faster access to auditable security guidance.", "Private inference for sensitive internal documents.", "Saved investigations improve analyst continuity.", "Version-linked feedback creates regression tests and measurable improvement."], 16)
    textbox(slide, 1.0, 6.52, 11.3, 0.38, "RATE → REVIEW EVIDENCE → ADD GOLDEN CASE → REFINE RETRIEVAL → REEVALUATE", 15, GREEN, True, align=PP_ALIGN.CENTER)

    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
